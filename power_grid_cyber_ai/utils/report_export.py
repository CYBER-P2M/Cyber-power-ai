"""Convert Markdown reports to formatted Word (.docx) documents."""
from __future__ import annotations
import io, re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

def _add_runs(paragraph, text):
    tokens = re.split(r"(\*\*.+?\*\*|\*.+?\*)", text)
    for token in tokens:
        if not token: continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2]); run.bold = True
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1]); run.italic = True
        else: paragraph.add_run(token)

def markdown_to_docx(markdown_text, title="PowerGrid Cyber AI Report"):
    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    h = document.add_heading(title, level=0); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lines = markdown_text.replace("\r\n", "\n").split("\n")
    in_code=False; code=[]
    for raw in lines:
        line=raw.rstrip()
        if line.strip().startswith("```"):
            if in_code:
                p=document.add_paragraph(); r=p.add_run("\n".join(code)); r.font.name="Consolas"; r.font.size=Pt(9); code=[]
            in_code=not in_code; continue
        if in_code: code.append(line); continue
        if not line.strip(): document.add_paragraph(); continue
        if line.startswith("|") and line.endswith("|"):
            cells=[c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells): continue
            t=document.add_table(rows=1, cols=len(cells)); t.style="Table Grid"
            for i,c in enumerate(cells): t.rows[0].cells[i].text=c
            continue
        m=re.match(r"^(#{1,6})\s+(.*)$", line)
        if m: document.add_heading(m.group(2).strip(), level=min(4,len(m.group(1)))); continue
        m=re.match(r"^\s*[-*+]\s+(.*)$", line)
        if m: p=document.add_paragraph(style="List Bullet"); _add_runs(p,m.group(1)); continue
        m=re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m: p=document.add_paragraph(style="List Number"); _add_runs(p,m.group(1)); continue
        p=document.add_paragraph(); _add_runs(p,line)
    b=io.BytesIO(); document.save(b); return b.getvalue()
